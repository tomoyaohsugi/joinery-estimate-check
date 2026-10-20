import io, re
from datetime import datetime
import streamlit as st
import pandas as pd
import numpy as np
import cv2
import pytesseract
from pytesseract import Output
from PIL import Image
from pypdf import PdfReader
from pdf2image import convert_from_bytes
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="建具工事 見積前チェック", page_icon="🏗️", layout="wide")

SYMBOL_RE = re.compile(r"(?<![A-Z0-9])((?:AW|AD|SD|WD|SW|AG|SG|FW|FD|WW|PG|PW|D|W)[-－_ ]?\d{1,3})(?!\d)", re.I)
SIZE_RES = [
    re.compile(r"(?:W|Ｗ)\s*[:=]?\s*(\d{2,5})\s*[×xX＊*]\s*(?:H|Ｈ)?\s*[:=]?\s*(\d{2,5})", re.I),
    re.compile(r"(\d{2,5})\s*[×xX＊*]\s*(\d{2,5})"),
    re.compile(r"(?:W|Ｗ)\s*[:=]?\s*(\d{2,5}).{0,20}?(?:H|Ｈ)\s*[:=]?\s*(\d{2,5})", re.I),
]
QTY_RE = re.compile(r"(?:数量|箇所|ヶ所|個数|台数|枚数|QTY)\s*[:：=]?\s*(\d{1,3})", re.I)
KEYWORDS = {
 "開閉形式": ["FIX","引違","片引","両引","片開","両開","縦すべり","横すべり","内倒","上げ下げ","排煙"],
 "材質": ["アルミ","スチール","ステンレス","木製","樹脂"],
 "ガラス": ["Low-E","複層","合わせガラス","強化ガラス","網入","型板","防火ガラス","透明ガラス","フロート"],
 "防火・性能": ["特定防火設備","防火設備","耐風圧","気密","水密","遮音","断熱"],
 "付属品": ["網戸","面格子","手すり","フック棒","オペレーター","ドアクローザー","フロアヒンジ","戸当り","シリンダー","錠","ガラリ","水切り","額縁"],
 "工事関連": ["撤去","処分","カバー工法","はつり","シール","コーキング","養生","搬入","揚重","足場","溶接","アンカー","自動ドア","排煙装置"],
}


def clean_symbol(s):
    return re.sub(r"[-－_ ]", "-", s.upper()).strip("-")


def kw(text, words):
    low=text.lower(); return "、".join([w for w in words if w.lower() in low])


def get_size(text):
    for rx in SIZE_RES:
        m=rx.search(text)
        if m: return int(m.group(1)), int(m.group(2))
    return None, None


def preprocess(img):
    arr=np.array(img.convert("RGB")); gray=cv2.cvtColor(arr,cv2.COLOR_RGB2GRAY)
    gray=cv2.resize(gray,None,fx=1.35,fy=1.35,interpolation=cv2.INTER_CUBIC)
    gray=cv2.bilateralFilter(gray,5,40,40)
    binary=cv2.adaptiveThreshold(gray,255,cv2.ADAPTIVE_THRESH_GAUSSIAN_C,cv2.THRESH_BINARY,31,11)
    return Image.fromarray(binary)


def ocr_image(img):
    data=pytesseract.image_to_data(preprocess(img),lang="jpn+eng",config="--oem 3 --psm 11 -c preserve_interword_spaces=1",output_type=Output.DICT)
    words=[]; confs=[]
    for text,conf in zip(data["text"],data["conf"]):
        text=str(text).strip()
        try: conf=float(conf)
        except: conf=-1
        if text and conf>=20:
            words.append(text); confs.append(conf)
    return " ".join(words), round(sum(confs)/len(confs),1) if confs else 0


def read_pdf(upload,dpi,force_ocr):
    raw=upload.getvalue(); reader=PdfReader(io.BytesIO(raw)); pages=[]
    progress=st.progress(0,text=f"{upload.name} を解析中")
    for i,page in enumerate(reader.pages,1):
        try: embedded=page.extract_text() or ""
        except: embedded=""
        method="PDF文字"; text=embedded; conf=100.0; error=""
        if force_ocr or len(embedded.strip())<20:
            try:
                image=convert_from_bytes(raw,dpi=dpi,first_page=i,last_page=i,fmt="png",thread_count=1)[0]
                text,conf=ocr_image(image); method="OCR"
            except Exception as e:
                method="OCR失敗"; error=f"{type(e).__name__}: {e}"
        pages.append({"page":i,"text":text,"method":method,"confidence":conf,"error":error})
        progress.progress(i/len(reader.pages),text=f"{upload.name}: {i}/{len(reader.pages)}ページ")
    progress.empty(); return pages


def extract(file_name,pages):
    obs=[]
    for p in pages:
        text=p["text"]
        for m in SYMBOL_RE.finditer(text):
            symbol=clean_symbol(m.group(1)); a=max(0,m.start()-100); b=min(len(text),m.end()+300)
            ctx=re.sub(r"\s+"," ",text[a:b]).strip(); w,h=get_size(ctx); qm=QTY_RE.search(ctx)
            obs.append({
              "建具符号":symbol,"W":w,"H":h,"数量":int(qm.group(1)) if qm else None,
              "開閉形式":kw(ctx,KEYWORDS["開閉形式"]),"材質":kw(ctx,KEYWORDS["材質"]),
              "ガラス":kw(ctx,KEYWORDS["ガラス"]),"防火・性能":kw(ctx,KEYWORDS["防火・性能"]),
              "付属品":kw(ctx,KEYWORDS["付属品"]),"工事関連":kw(ctx,KEYWORDS["工事関連"]),
              "参照ファイル":file_name,"参照ページ":p["page"],"読取方法":p["method"],
              "OCR信頼度":p["confidence"],"根拠抜粋":ctx[:350]
            })
    if not obs: return pd.DataFrame(),[]
    conflicts=[]; rows=[]
    for symbol,group in pd.DataFrame(obs).groupby("建具符号"):
        base=group.iloc[0].to_dict(); base["参照ページ"]=", ".join(map(str,sorted(group["参照ページ"].unique())))
        base["根拠抜粋"]=" | ".join(group["根拠抜粋"].drop_duplicates().head(3))
        for f in ["W","H","数量","開閉形式","材質","ガラス","防火・性能","付属品","工事関連"]:
            vals=[str(v) for v in group[f].dropna().unique() if str(v).strip()]
            if len(vals)==1: base[f]=int(float(vals[0])) if f in ["W","H","数量"] else vals[0]
            elif len(vals)>1:
                base[f]=" / ".join(vals); conflicts.append({"建具符号":symbol,"項目":f,"差異内容":base[f],"参照ページ":base["参照ページ"]})
        rows.append(base)
    return pd.DataFrame(rows),conflicts


def judge(df,conflicts):
    conflict_symbols={c["建具符号"] for c in conflicts}; statuses=[]; reasons=[]
    for _,r in df.iterrows():
        critical=[x for x in ["W","H","数量","開閉形式"] if pd.isna(r.get(x)) or str(r.get(x)).strip() in ["","nan"]]
        caution=[x for x in ["ガラス","防火・性能","付属品"] if pd.isna(r.get(x)) or not str(r.get(x)).strip()]
        why=[]
        if r["建具符号"] in conflict_symbols: why.append("図面内差異あり")
        if critical: why.append("重要情報不足: "+"、".join(critical))
        if caution: why.append("確認推奨: "+"、".join(caution))
        status="見積前確認必須" if critical or r["建具符号"] in conflict_symbols else "要確認" if caution else "見積準備OK"
        statuses.append(status); reasons.append(" / ".join(why))
    df["判定"]=statuses; df["要確認理由"]=reasons; return df


def excel_bytes(df,conflicts,project):
    wb=Workbook(); ws=wb.active; ws.title="建具拾い出し一覧"
    sheets=[(ws,df),(wb.create_sheet("見積前要確認"),df[df["判定"]!="見積準備OK"]),(wb.create_sheet("図面不整合"),pd.DataFrame(conflicts))]
    for sh,data in sheets:
        if data.empty: sh["A1"]="該当なし"; continue
        for c,h in enumerate(data.columns,1):
            x=sh.cell(1,c,h); x.font=Font(bold=True,color="FFFFFF"); x.fill=PatternFill("solid",fgColor="17365D")
        for r,row in enumerate(data.itertuples(index=False),2):
            for c,v in enumerate(row,1): sh.cell(r,c,v).alignment=Alignment(wrap_text=True,vertical="top")
        for c in range(1,sh.max_column+1): sh.column_dimensions[get_column_letter(c)].width=18
        sh.freeze_panes="A2"; sh.auto_filter.ref=sh.dimensions
    info=wb.create_sheet("案件情報"); info.append(["案件名",project]); info.append(["出力日時",datetime.now().strftime("%Y-%m-%d %H:%M")]); info.append(["注意","見積前確認支援用です。元図面の確認を優先してください。"]) 
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue()

st.title("🏗️ 建具工事 見積前チェック・図面解析システム")
st.caption("画像PDFを日本語OCRし、建具符号・寸法・仕様候補・不足情報を整理します。")
with st.sidebar:
    project=st.text_input("案件名",placeholder="例：○○新築工事")
    dpi=st.select_slider("OCR解像度（DPI）",[200,250,300],250)
    force=st.checkbox("全ページをOCRする",False)
    st.info("初回は250DPIを推奨。読取が悪い場合のみ300DPIを試してください。")
files=st.file_uploader("建築図面PDFをアップロード",type=["pdf"],accept_multiple_files=True)
if st.button("解析開始",type="primary",disabled=not files,use_container_width=True):
    dfs=[]; conflicts=[]; warnings=[]
    for f in files:
        pages=read_pdf(f,dpi,force)
        failed=[p["page"] for p in pages if p["method"]=="OCR失敗"]
        if failed: warnings.append(f"{f.name}: OCR失敗ページ {failed}")
        df,cc=extract(f.name,pages)
        if not df.empty: dfs.append(df)
        conflicts.extend(cc)
    if not dfs:
        st.error("OCR後も建具符号を抽出できませんでした。300DPI、全ページOCR、または図面分割を試してください。")
    else:
        result=judge(pd.concat(dfs,ignore_index=True),conflicts)
        st.session_state.update(result=result,conflicts=conflicts,warnings=warnings,project=project or "未設定案件")

if "result" in st.session_state:
    df=st.session_state.result; counts=df["判定"].value_counts(); cols=st.columns(4)
    cols[0].metric("抽出建具数",len(df)); cols[1].metric("見積準備OK",int(counts.get("見積準備OK",0))); cols[2].metric("要確認",int(counts.get("要確認",0))); cols[3].metric("確認必須",int(counts.get("見積前確認必須",0)))
    for w in st.session_state.warnings: st.warning(w)
    t1,t2,t3,t4=st.tabs(["要確認一覧","建具拾い出し","図面不整合","Excel出力"])
    with t1: st.dataframe(df[df["判定"]!="見積準備OK"],use_container_width=True,hide_index=True)
    with t2: st.dataframe(df,use_container_width=True,hide_index=True)
    with t3:
        if st.session_state.conflicts: st.dataframe(pd.DataFrame(st.session_state.conflicts),use_container_width=True,hide_index=True)
        else: st.success("抽出範囲では差異なし")
    with t4:
        st.download_button("Excel報告書をダウンロード",excel_bytes(df,st.session_state.conflicts,st.session_state.project),file_name=f"{st.session_state.project}_建具見積前チェック.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)

st.divider(); st.warning("OCR結果は補助情報です。細い文字、回転文字、表構造、図形数量は誤認識する場合があります。見積確定前に元図面を確認してください。")
